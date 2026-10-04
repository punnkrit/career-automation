from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt


def markdown_to_docx(markdown_path: Path, docx_path: Path) -> Path:
    text = markdown_path.read_text(encoding="utf-8")
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line:
            document.add_paragraph()
        elif line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
        elif line.startswith("- "):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(docx_path)
    return docx_path
